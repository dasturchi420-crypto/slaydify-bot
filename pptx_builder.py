"""AI bergan JSON dan chiroyli .pptx fayl yasaydi (16:9)."""
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

W, H = Inches(13.333), Inches(7.5)

THEMES = {
    "ocean": {"name": "🌊 Okean", "bg": "0B2545", "panel": "13315C", "title": "FFFFFF",
              "text": "DCE8F5", "accent": "4CC9F0", "font": "Calibri"},
    "light": {"name": "☀️ Oq-minimal", "bg": "FFFFFF", "panel": "F1F4F9", "title": "1B2A41",
              "text": "324A5F", "accent": "2563EB", "font": "Calibri"},
    "emerald": {"name": "🌿 Zumrad", "bg": "0F2E26", "panel": "164A3C", "title": "FFFFFF",
                "text": "D7EFE6", "accent": "34D399", "font": "Calibri"},
    "sunset": {"name": "🌅 Quyosh", "bg": "2B1B3D", "panel": "412A5C", "title": "FFFFFF",
               "text": "F1E4FF", "accent": "FF8A5B", "font": "Calibri"},
    "mono": {"name": "⚫ Qora", "bg": "111111", "panel": "1E1E1E", "title": "FFFFFF",
             "text": "D4D4D4", "accent": "FACC15", "font": "Calibri"},
    "sand": {"name": "🏜 Qum", "bg": "FAF5EB", "panel": "F0E6D2", "title": "3D2C1E",
             "text": "5A4636", "accent": "C2410C", "font": "Georgia"},
}


def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h)


def _rect(slide, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(color)
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def _text(slide, x, y, w, h, text, size, color, font, bold=False, align=PP_ALIGN.LEFT,
          anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size, r.font.bold, r.font.name = Pt(size), bold, font
    r.font.color.rgb = rgb(color)
    return tb


def _background(slide, color):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = rgb(color)


def build(data: dict, theme_key: str, path: str, lang: str = "uz") -> str:
    t = THEMES.get(theme_key, THEMES["ocean"])
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]

    # --- Sarlavha slaydi ---
    s = prs.slides.add_slide(blank)
    _background(s, t["bg"])
    _rect(s, 0, 0, Inches(0.35), H, t["accent"])
    _rect(s, Inches(8.9), Inches(-1.5), Inches(6.5), Inches(6.5), t["panel"], MSO_SHAPE.OVAL)
    _rect(s, Inches(10.6), Inches(4.6), Inches(3.8), Inches(3.8), t["accent"], MSO_SHAPE.OVAL)
    _text(s, Inches(1.0), Inches(2.1), Inches(9.2), Inches(2.4), data.get("title", ""), 46,
          t["title"], t["font"], bold=True, anchor=MSO_ANCHOR.BOTTOM)
    _rect(s, Inches(1.05), Inches(4.7), Inches(1.6), Inches(0.08), t["accent"])
    _text(s, Inches(1.0), Inches(4.95), Inches(8.5), Inches(1.2), data.get("subtitle", ""), 22,
          t["text"], t["font"])

    # --- Kontent slaydlari ---
    slides = data["slides"]
    total = len(slides) + 1
    for i, item in enumerate(slides, start=2):
        s = prs.slides.add_slide(blank)
        _background(s, t["bg"])
        _rect(s, 0, 0, W, Inches(1.5), t["panel"])
        _rect(s, 0, 0, Inches(0.35), Inches(1.5), t["accent"])
        _text(s, Inches(0.9), Inches(0.2), Inches(11.5), Inches(1.1), item["title"], 32,
              t["title"], t["font"], bold=True, anchor=MSO_ANCHOR.MIDDLE)

        bullets = item["bullets"][:5]
        top, gap = Inches(1.95), Inches(0.98)
        size = 22 if max(len(b) for b in bullets) < 110 else 19
        for j, b in enumerate(bullets):
            y = top + gap * j
            _rect(s, Inches(0.9), y + Inches(0.14), Inches(0.28), Inches(0.28), t["accent"],
                  MSO_SHAPE.OVAL)
            _text(s, Inches(1.45), y, Inches(11.2), Inches(0.9), b, size, t["text"], t["font"],
                  anchor=MSO_ANCHOR.TOP)

        _text(s, Inches(11.6), Inches(6.95), Inches(1.3), Inches(0.4), f"{i} / {total}", 12,
              t["accent"], t["font"], align=PP_ALIGN.RIGHT)
        if item.get("notes"):
            s.notes_slide.notes_text_frame.text = item["notes"]

    # --- Yakuniy slayd ---
    thanks = {"uz": "E'tiboringiz uchun rahmat!", "ru": "Спасибо за внимание!",
              "en": "Thank you for your attention!"}.get(lang, "Thank you!")
    s = prs.slides.add_slide(blank)
    _background(s, t["bg"])
    _rect(s, Inches(-1.5), Inches(4.2), Inches(5), Inches(5), t["panel"], MSO_SHAPE.OVAL)
    _rect(s, Inches(10.8), Inches(-1.2), Inches(4), Inches(4), t["accent"], MSO_SHAPE.OVAL)
    _text(s, Inches(1.0), Inches(2.6), Inches(11.3), Inches(2.0), thanks, 48, t["title"],
          t["font"], bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    _rect(s, Inches(5.9), Inches(4.7), Inches(1.5), Inches(0.08), t["accent"])

    prs.save(path)
    return path
